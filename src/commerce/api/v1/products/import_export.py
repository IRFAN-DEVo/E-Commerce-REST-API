"""
DjangoMart — Product Import / Export
=====================================
Supports:
  • CSV import / export
  • Excel (.xlsx) import / export
  • Optional ZIP archive with product images (main + gallery)
  • Skip / Update / Replace modes
  • Automatic Category and Brand creation
  • Detailed import statistics and error reporting

ZIP image import is fully optional.  If no ZIP is provided the importer
behaves exactly as before (backward compatible).
"""

import csv
import io
import time
import zipfile
from decimal import Decimal, InvalidOperation

from django.core.files.base import ContentFile
from django.db import transaction
from django.http import HttpResponse

import openpyxl
from rest_framework import status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from products.models import Brand, Category, Product, ProductImage
from .permissions import IsOwner

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

ALLOWED_IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}
MAX_IMAGE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB per image
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024  # 100 MB max uncompressed
MAX_ARCHIVE_FILES = 1000


# ===========================================================================
# Export Views (unchanged)
# ===========================================================================


class ProductExportCSVView(APIView):
    permission_classes = [IsAuthenticated, IsOwner]

    def get(self, request):
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="products_export.csv"'

        writer = csv.writer(response)
        writer.writerow([
            "Product Name", "SKU", "Description", "Category", "Brand",
            "Price", "Stock", "Is Active",
        ])

        products = Product.objects.select_related("category", "brand").all()
        for product in products:
            writer.writerow([
                product.name,
                product.sku or "",
                product.description,
                product.category.name if product.category is not None else "",
                product.brand.name if product.brand is not None else "",
                product.price,
                product.stock,
                "Yes" if product.is_available else "No",
            ])

        return response


class ProductExportExcelView(APIView):
    permission_classes = [IsAuthenticated, IsOwner]

    def get(self, request):
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="products_export.xlsx"'

        wb = openpyxl.Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet(title="Products")
        else:
            ws.title = "Products"

        headers = [
            "Product Name", "SKU", "Description", "Category", "Brand",
            "Price", "Stock", "Is Active",
        ]
        ws.append(headers)

        products = Product.objects.select_related("category", "brand").all()
        for product in products:
            ws.append([
                product.name,
                product.sku or "",
                product.description,
                product.category.name if product.category is not None else "",
                product.brand.name if product.brand is not None else "",
                float(product.price) if product.price is not None else 0.0,
                product.stock,
                "Yes" if product.is_available else "No",
            ])

        wb.save(response)
        return response


# ===========================================================================
# Template Download (extended with SKU & image columns)
# ===========================================================================


class ProductTemplateView(APIView):
    """
    Download a sample CSV template that includes SKU and image columns.
    """
    permission_classes = [IsAuthenticated, IsOwner]

    def get(self, request):
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="product_import_template.csv"'

        writer = csv.writer(response)
        writer.writerow([
            "Product Name", "SKU", "Description", "Category", "Brand",
            "Price", "Stock", "Is Active",
            "Main Image", "Gallery Images",
        ])
        # Sample rows
        writer.writerow([
            "iPhone 16", "IPH-16-BLK", "Latest Apple smartphone", "Phones", "Apple",
            "999.99", "20", "Yes",
            "iphone.jpg", "iphone1.jpg,iphone2.jpg",
        ])
        writer.writerow([
            "Galaxy S25", "SAM-S25-SLV", "Samsung flagship", "Phones", "Samsung",
            "899.99", "15", "Yes",
            "galaxy.jpg", "galaxy1.jpg,galaxy2.jpg",
        ])
        writer.writerow([
            "Sample Laptop", "DEL-XPS-15", "A sample laptop description", "Computers", "Dell",
            "1299.99", "8", "Yes",
            "laptop.jpg", "",
        ])

        return response


# ===========================================================================
# Helpers
# ===========================================================================


def _normalise_key(raw_path: str) -> str:
    """
    Convert any ZIP-internal path to a normalised lookup key:
      images/photo.jpg  → photo.jpg
      .\\photo.jpg       → photo.jpg
      PHOTO.JPG          → photo.jpg
    """
    basename = raw_path.replace("\\", "/").split("/")[-1]
    return basename.strip().lower()


def _closest_matches(name: str, keys: list[str], n: int = 3) -> list[str]:
    """
    Return up to n keys that share the most characters with `name`.
    Uses a simple character-overlap score (no external libs required).
    """
    target = name.lower()
    scored = []
    for k in keys:
        common = sum(c in k for c in target)
        scored.append((common, k))
    scored.sort(key=lambda x: -x[0])
    return [k for _, k in scored[:n]]


def _validate_and_load_zip(zip_file_obj):
    """
    Validate the uploaded ZIP file and build a normalised filename → bytes mapping.
    Enforces security against path traversal (Zip Slip), absolute paths, zip bombs,
    and invalid file types.
    """
    errors: list[str] = []

    raw_bytes = zip_file_obj.read()
    if not zipfile.is_zipfile(io.BytesIO(raw_bytes)):
        raise ValueError("The uploaded file is not a valid ZIP archive.")

    image_lookup: dict[str, bytes] = {}
    key_to_original: dict[str, str] = {}
    total_uncompressed = 0

    with zipfile.ZipFile(io.BytesIO(raw_bytes), "r") as zf:
        infolist = zf.infolist()
        if len(infolist) > MAX_ARCHIVE_FILES:
            raise ValueError(f"ZIP archive contains too many files ({len(infolist)}). Limit is {MAX_ARCHIVE_FILES}.")

        for info in infolist:
            if info.is_dir():
                continue

            total_uncompressed += info.file_size
            if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ValueError("ZIP archive exceeds maximum uncompressed size limit (100 MB).")

            raw_path = info.filename
            normalised = raw_path.replace("\\", "/")

            # --- Security: block path-traversal & absolute paths ---
            if normalised.startswith("/") or ":" in normalised[:4] or any(part == ".." for part in normalised.split("/")):
                errors.append(
                    f"ZIP security: '{raw_path}' contains path traversal or absolute path and was skipped."
                )
                continue

            key = _normalise_key(raw_path)

            # --- Validate extension ---
            ext = key.rsplit(".", 1)[-1] if "." in key else ""
            if ext not in ALLOWED_IMAGE_EXTENSIONS:
                if ext:
                    errors.append(
                        f"ZIP: '{raw_path}' skipped — unsupported extension (.{ext}). "
                        f"Allowed: {', '.join(sorted(ALLOWED_IMAGE_EXTENSIONS))}."
                    )
                continue

            # --- Validate individual size ---
            if info.file_size > MAX_IMAGE_SIZE_BYTES:
                errors.append(
                    f"ZIP: '{raw_path}' exceeds the 10 MB size limit and was skipped."
                )
                continue

            # --- Warn on duplicate basenames ---
            if key in image_lookup:
                errors.append(
                    f"ZIP: Duplicate basename '{key}' (from '{raw_path}' and '{key_to_original[key]}') — "
                    "only the last occurrence will be used."
                )

            image_lookup[key] = zf.read(raw_path)
            key_to_original[key] = raw_path

    return image_lookup, errors


def _lookup_image(name: str, image_lookup: dict[str, bytes]) -> tuple[bytes | None, str]:
    key = _normalise_key(name)
    data = image_lookup.get(key)
    return data, key


def _assign_main_image(
    product: Product,
    filename: str,
    image_lookup: dict[str, bytes],
    all_keys: list[str],
) -> tuple[bool, str | None]:
    data, key = _lookup_image(filename, image_lookup)
    if data is None:
        hints = _closest_matches(key, all_keys)
        hint_str = ", ".join(hints) if hints else "(no similar filenames found)"
        return False, (
            f"Main image '{filename}' not found in ZIP. Nearest matches: {hint_str}"
        )
    product.image.save(key, ContentFile(data), save=True)
    return True, None


def _assign_gallery_images(
    product: Product,
    filenames: list[str],
    image_lookup: dict[str, bytes],
    all_keys: list[str],
    replace: bool = False,
) -> tuple[int, list[str]]:
    if replace:
        product.images.all().delete()

    saved = 0
    warnings: list[str] = []

    for raw_name in filenames:
        raw_name = raw_name.strip()
        if not raw_name:
            continue
        data, key = _lookup_image(raw_name, image_lookup)
        if data is None:
            hints = _closest_matches(key, all_keys)
            hint_str = ", ".join(hints) if hints else "(no similar filenames found)"
            warnings.append(
                f"Gallery image '{raw_name}' not found in ZIP. Nearest matches: {hint_str}"
            )
            continue
        pi = ProductImage(product=product, alt_text=key)
        pi.image.save(key, ContentFile(data), save=True)
        saved += 1

    return saved, warnings


# ===========================================================================
# Import View (with Preview & Dry Run Support)
# ===========================================================================


class ProductImportView(APIView):
    """
    Import products from a CSV or Excel file.
    Optionally accepts images_zip file for images.
    Supports dry_run / validate_only preview mode.
    """

    permission_classes = [IsAuthenticated, IsOwner]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        start_time = time.time()

        file_obj = request.FILES.get("file")
        zip_obj = request.FILES.get("images_zip")
        import_mode = request.data.get("import_mode", "skip").strip().lower()  # skip | update | replace
        dry_run_param = request.data.get("dry_run") or request.data.get("validate_only") or request.query_params.get("dry_run")
        is_dry_run = str(dry_run_param).strip().lower() in ("true", "1", "yes")

        # ------------------------------------------------------------------
        # 1. Validate product file
        # ------------------------------------------------------------------
        if not file_obj:
            return Response(
                {"error": "No product data file uploaded."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        filename = file_obj.name
        rows: list[dict] = []

        try:
            if filename.endswith(".csv"):
                decoded_file = file_obj.read().decode("utf-8-sig")  # handle BOM if present
                io_string = io.StringIO(decoded_file)
                reader = csv.DictReader(io_string)
                for row in reader:
                    rows.append(row)

            elif filename.endswith(".xlsx"):
                wb = openpyxl.load_workbook(file_obj, data_only=True)
                ws = wb.active
                if ws is None:
                    if wb.worksheets:
                        ws = wb.worksheets[0]
                    else:
                        raise ValueError("The uploaded Excel file contains no worksheets.")

                if ws.max_row < 1:
                    raise ValueError("The uploaded Excel file is empty.")

                headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if any(cell is not None for cell in row):
                        rows.append(dict(zip(headers, row)))
            else:
                return Response(
                    {"error": "Unsupported file format. Please upload CSV or Excel (.xlsx)."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        except Exception as exc:
            return Response(
                {"error": f"Error parsing file: {exc}"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not rows:
            return Response(
                {"error": "The uploaded file does not contain any data rows."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ------------------------------------------------------------------
        # 2. Validate and load ZIP (optional)
        # ------------------------------------------------------------------
        image_lookup: dict[str, bytes] = {}
        zip_errors: list[str] = []
        has_zip = zip_obj is not None

        if has_zip:
            try:
                image_lookup, zip_errors = _validate_and_load_zip(zip_obj)
            except ValueError as exc:
                return Response(
                    {"error": str(exc)},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        all_zip_keys = list(image_lookup.keys())

        # ------------------------------------------------------------------
        # 3. Pre-load reference data
        # ------------------------------------------------------------------
        category_map = {c.name.lower(): c for c in Category.objects.all()}
        brand_map = {b.name.lower(): b for b in Brand.objects.all()}

        # Map existing products by SKU (primary identifier) and Name
        products_by_sku: dict[str, Product] = {
            p.sku.lower(): p
            for p in Product.objects.select_related("category", "brand").exclude(sku__isnull=True).exclude(sku="")
        }
        products_by_name: dict[str, Product] = {
            p.name.lower(): p
            for p in Product.objects.select_related("category", "brand").all()
        }

        # ------------------------------------------------------------------
        # 4. Process & validate rows
        # ------------------------------------------------------------------
        imported = 0
        updated = 0
        skipped = 0
        failed = 0
        main_images_imported = 0
        gallery_images_imported = 0

        total_images_found = 0
        total_images_missing = 0

        errors: list[str] = list(zip_errors)
        failed_rows: list[dict] = []
        preview_rows: list[dict] = []

        seen_skus: set[str] = set()

        for index, row in enumerate(rows, start=2):
            row_errors: list[str] = []

            # Normalize row keys
            row_data = {
                str(k).strip().lower(): v
                for k, v in row.items()
                if k is not None
            }

            name = (
                row_data.get("product name")
                or row_data.get("name")
            )
            sku = str(row_data.get("sku") or "").strip()
            description = str(row_data.get("description") or "").strip()
            category_name = row_data.get("category")
            brand_name = row_data.get("brand")
            price_val = row_data.get("price")
            stock_val = row_data.get("stock", 0)
            is_active_val = row_data.get("is active", "Yes")

            # Image fields
            main_image_name = str(row_data.get("main image") or "").strip()
            gallery_raw = str(row_data.get("gallery images") or "").strip()
            sep = "|" if "|" in gallery_raw else ","
            gallery_names = [g.strip() for g in gallery_raw.split(sep) if g.strip()]

            # --- Required field check ---
            if not name:
                row_errors.append("Product name is required.")
            if not category_name:
                row_errors.append("Category is required.")
            if not brand_name:
                row_errors.append("Brand is required.")
            if price_val is None or str(price_val).strip() == "":
                row_errors.append("Price is required.")

            name_clean = str(name).strip() if name else ""

            # --- Duplicate SKU in CSV check ---
            if sku:
                sku_clean = sku.strip()
                if sku_clean.lower() in seen_skus:
                    row_errors.append(f"Duplicate SKU '{sku_clean}' found within the import file.")
                else:
                    seen_skus.add(sku_clean.lower())
            else:
                sku_clean = None

            # --- Price validation ---
            price = None
            if price_val is not None and str(price_val).strip() != "":
                try:
                    price = Decimal(str(price_val).strip())
                    if price <= 0:
                        row_errors.append("Price must be greater than zero.")
                except (InvalidOperation, ValueError):
                    row_errors.append("Invalid price format.")

            # --- Stock validation ---
            stock = 0
            if stock_val is not None and str(stock_val).strip() != "":
                try:
                    stock = int(float(str(stock_val).strip()))
                    if stock < 0:
                        row_errors.append("Stock quantity cannot be negative.")
                except (ValueError, TypeError):
                    row_errors.append("Invalid stock format.")

            is_active = str(is_active_val).strip().lower() in ["yes", "true", "1", "y"]
            stock_status = "IN_STOCK" if stock > 0 else "OUT_OF_STOCK"

            # --- Image check ---
            img_found_count = 0
            img_missing_count = 0
            img_status = "None"

            if has_zip:
                if main_image_name:
                    data, _ = _lookup_image(main_image_name, image_lookup)
                    if data:
                        img_found_count += 1
                    else:
                        img_missing_count += 1
                        row_errors.append(f"Main image '{main_image_name}' not found in ZIP.")

                for gname in gallery_names:
                    data, _ = _lookup_image(gname, image_lookup)
                    if data:
                        img_found_count += 1
                    else:
                        img_missing_count += 1
                        row_errors.append(f"Gallery image '{gname}' not found in ZIP.")

                total_images_found += img_found_count
                total_images_missing += img_missing_count

                if (main_image_name or gallery_names):
                    if img_missing_count == 0:
                        img_status = f"All {img_found_count} images found"
                    elif img_found_count > 0:
                        img_status = f"{img_found_count} found, {img_missing_count} missing"
                    else:
                        img_status = f"All {img_missing_count} images missing"
            else:
                if main_image_name or gallery_names:
                    img_status = "Images referenced (no ZIP uploaded)"

            # --- Match existing product ---
            matched_product = None
            if sku_clean and sku_clean.lower() in products_by_sku:
                matched_product = products_by_sku[sku_clean.lower()]
            elif name_clean and name_clean.lower() in products_by_name:
                matched_product = products_by_name[name_clean.lower()]

            # Determine action
            if row_errors:
                action = "ERROR"
            elif matched_product:
                if import_mode == "skip":
                    action = "SKIP"
                else:
                    action = "UPDATE"
            else:
                action = "CREATE"

            # -------------------------------------------------------
            # If DRY RUN: record row preview and skip DB persistence
            # -------------------------------------------------------
            if is_dry_run:
                preview_rows.append({
                    "row": index,
                    "name": name_clean,
                    "sku": sku_clean or "—",
                    "category": str(category_name or "").strip(),
                    "brand": str(brand_name or "").strip(),
                    "price": str(price) if price is not None else str(price_val or ""),
                    "stock": stock,
                    "action": action,
                    "image_status": img_status,
                    "errors": row_errors,
                })
                continue

            # -------------------------------------------------------
            # ACTUAL IMPORT EXECUTION
            # -------------------------------------------------------
            if row_errors:
                failed += 1
                reason = "; ".join(row_errors)
                errors.append(f"Row {index}: {reason}")
                failed_rows.append({
                    "row": index,
                    "name": name_clean,
                    "sku": sku_clean or "",
                    "reason": reason,
                })
                continue

            try:
                with transaction.atomic():
                    # Category resolution
                    cat_key = str(category_name).strip().lower()
                    if cat_key not in category_map:
                        new_cat = Category.objects.create(name=str(category_name).strip())
                        category_map[cat_key] = new_cat
                    category = category_map[cat_key]

                    # Brand resolution
                    brand_key = str(brand_name).strip().lower()
                    if brand_key not in brand_map:
                        new_brand = Brand.objects.create(name=str(brand_name).strip())
                        brand_map[brand_key] = new_brand
                    brand = brand_map[brand_key]

                    if matched_product:
                        if import_mode == "skip":
                            skipped += 1
                            continue

                        product = matched_product
                        product.name = name_clean
                        if sku_clean:
                            product.sku = sku_clean
                        product.description = description if description else product.description
                        product.category = category
                        product.brand = brand
                        product.price = price
                        product.stock = stock
                        product.stock_status = stock_status
                        product.is_available = is_active

                        if import_mode == "replace" and has_zip:
                            product.images.all().delete()

                        product.save()
                        updated += 1
                    else:
                        product = Product.objects.create(
                            name=name_clean,
                            sku=sku_clean,
                            description=description,
                            category=category,
                            brand=brand,
                            price=price,
                            stock=stock,
                            stock_status=stock_status,
                            is_available=is_active,
                        )
                        if sku_clean:
                            products_by_sku[sku_clean.lower()] = product
                        products_by_name[name_clean.lower()] = product
                        imported += 1

                    # Image assignment
                    if has_zip:
                        if main_image_name:
                            ok, warn = _assign_main_image(
                                product, main_image_name, image_lookup, all_zip_keys
                            )
                            if ok:
                                main_images_imported += 1
                            elif warn:
                                errors.append(f"Row {index}: {warn}")

                        if gallery_names:
                            replace_gallery = import_mode == "replace"
                            count, warns = _assign_gallery_images(
                                product,
                                gallery_names,
                                image_lookup,
                                all_zip_keys,
                                replace=replace_gallery,
                            )
                            gallery_images_imported += count
                            for w in warns:
                                errors.append(f"Row {index}: {w}")

            except Exception as exc:
                failed += 1
                reason = f"Database error: {exc}"
                errors.append(f"Row {index}: {reason}")
                failed_rows.append({
                    "row": index,
                    "name": name_clean,
                    "sku": sku_clean or "",
                    "reason": reason,
                })

        # ------------------------------------------------------------------
        # 5. Response formatting
        # ------------------------------------------------------------------
        execution_time = round(time.time() - start_time, 2)

        if is_dry_run:
            total_rows = len(preview_rows)
            valid_rows = sum(1 for r in preview_rows if len(r["errors"]) == 0)
            invalid_rows = total_rows - valid_rows
            new_products = sum(1 for r in preview_rows if r["action"] == "CREATE")
            products_to_update = sum(1 for r in preview_rows if r["action"] == "UPDATE")

            return Response({
                "success": True,
                "dry_run": True,
                "summary": {
                    "total_rows": total_rows,
                    "valid_rows": valid_rows,
                    "invalid_rows": invalid_rows,
                    "new_products": new_products,
                    "products_to_update": products_to_update,
                    "images_found": total_images_found,
                    "images_missing": total_images_missing,
                },
                "preview": preview_rows,
                "zip_errors": zip_errors,
                "execution_time": execution_time,
            }, status=status.HTTP_200_OK)

        response_data: dict = {
            "success": True,
            "dry_run": False,
            "imported": imported,
            "updated": updated,
            "skipped": skipped,
            "failed": failed,
            "errors": errors,
            "failed_rows": failed_rows,
            "execution_time": execution_time,
        }

        if has_zip:
            response_data["main_images_imported"] = main_images_imported
            response_data["gallery_images_imported"] = gallery_images_imported

        return Response(response_data, status=status.HTTP_200_OK)
